"""PDF and Word uploads, with real files made in the test."""
import io

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter

from app.sources import SourceError, extract_text

LINES = ["Photosynthesis turns light energy into chemical energy.", "The Calvin cycle builds glucose."]


def make_pdf(lines: list[str]) -> bytes:
    """A minimal one page PDF with these lines as real text (no lines: a page with no text, like a scan)."""
    ops = ["BT /F1 12 Tf 72 720 Td"] + [f"({line}) Tj 0 -16 Td" for line in lines] + ["ET"]
    stream = "\n".join(ops).encode() if lines else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, obj in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf_text_is_extracted():
    text = extract_text("Notes.PDF", make_pdf(LINES), 10_000)
    assert "Photosynthesis turns light energy into chemical energy." in text
    assert "Calvin cycle builds glucose" in text


def test_scanned_pdf_without_text_is_refused_with_a_helpful_message():
    with pytest.raises(SourceError) as e:
        extract_text("scan.pdf", make_pdf([]), 10_000)
    assert e.value.status == 400 and "scanned" in str(e.value)


def test_password_protected_pdf_is_refused():
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf(LINES))))
    writer.encrypt(user_password="secret", algorithm="RC4-128")
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(SourceError) as e:
        extract_text("locked.pdf", buf.getvalue(), 10_000)
    assert "password" in str(e.value)


def test_damaged_files_are_refused_not_crashed():
    for name, data in (("bad.pdf", b"%PDF-1.4 this is not really a pdf"), ("bad.pdf", b"hello"),
                       ("bad.docx", b"not a zip"), ("bad.docx", b"PK\x03\x04broken")):
        with pytest.raises(SourceError) as e:
            extract_text(name, data, 10_000)
        assert e.value.status == 400 and "damaged" in str(e.value), name


def test_docx_paragraphs_and_tables_are_extracted():
    data = make_docx(LINES, table=[["Stage", "Where"], ["Light reactions", "Thylakoid"]])
    text = extract_text("notes.docx", data, 10_000)
    assert LINES[0] in text and LINES[1] in text
    assert "Light reactions | Thylakoid" in text


def test_old_doc_asks_for_docx_or_pdf():
    with pytest.raises(SourceError) as e:
        extract_text("notes.doc", b"\xd0\xcf\x11\xe0", 10_000)
    assert e.value.status == 400 and ".docx" in str(e.value)


def test_extracted_text_still_respects_the_character_limit():
    with pytest.raises(SourceError) as e:
        extract_text("long.docx", make_docx(["word " * 5000]), 1000)
    assert e.value.status == 413
    with pytest.raises(SourceError) as e:
        extract_text("empty.docx", make_docx([]), 1000)
    assert "empty" in str(e.value)
