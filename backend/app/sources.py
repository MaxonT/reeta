"""Source material: turn an upload into clean text, and combine a session's sources for the question writer.
To support another file type, add an extractor to _EXTRACTORS. Nothing else needs to change."""
import io
from pathlib import Path

PASTED = "pasted notes"
MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # a file; PDFs and Word files carry fonts and images. The text limit still applies.


class SourceError(ValueError):
    """An upload we refuse. `status` is the HTTP code the API should answer with."""

    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "latin-1"):  # latin-1 accepts any byte, so this never fails
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise AssertionError("unreachable")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise SourceError("That PDF is password protected. Remove the password and upload it again.")
        pages = ((page.extract_text() or "").strip() for page in reader.pages)
        text = "\n\n".join("\n".join(line.rstrip() for line in page.splitlines()) for page in pages)  # PDFs pad lines
    except SourceError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise SourceError("That PDF could not be read. Is the file damaged?") from exc
    if not text.strip():
        raise SourceError("That PDF has no text to read. A scanned PDF is only pictures of pages; "
                          "upload a PDF with selectable text, or paste the notes.")
    return text


def _docx_text(data: bytes) -> str:
    from zipfile import BadZipFile
    from docx import Document
    try:
        doc = Document(io.BytesIO(data))
    except (BadZipFile, KeyError, ValueError) as exc:
        raise SourceError("That Word file could not be read. Is the file damaged?") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # notes often live in tables
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _old_doc(data: bytes) -> str:
    raise SourceError("Old Word .doc files are not supported. In Word, use Save As > .docx or PDF, then upload that.")


_EXTRACTORS = {".txt": _decode_text, ".md": _decode_text, ".pdf": _pdf_text, ".docx": _docx_text, ".doc": _old_doc}
SUPPORTED = sorted(ext for ext, fn in _EXTRACTORS.items() if fn is not _old_doc)


def extract_text(filename: str, data: bytes, max_chars: int) -> str:
    """Bytes of an uploaded file -> clean text. Raises SourceError if empty, too long, or an unsupported type."""
    ext = Path(filename or "").suffix.lower()
    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        raise SourceError(f"Unsupported file type {ext or '(none)'}. Upload: {', '.join(SUPPORTED)}")
    return clean_text(extractor(data), max_chars)


def clean_text(text: str, max_chars: int) -> str:
    """Normalise pasted or extracted text. Raises SourceError if empty or too long."""
    text = text.replace("\r\n", "\n").replace("\x00", "").strip()
    if not text:
        raise SourceError("That source is empty.")
    if len(text) > max_chars:
        raise SourceError(f"That source is too long ({len(text):,} characters, the limit is {max_chars:,}).", 413)
    return text


def combine_sources(sources: list[dict], max_chars: int) -> str | None:
    """A session's sources as one block of notes for the question writer, cut to max_chars.
    None if there are no sources, so the caller can fall back to the default notes."""
    parts = [f"## {s['filename']}\n{s['text']}" for s in sources]
    if not parts:
        return None
    return "\n\n".join(parts)[:max_chars]
